from pathlib import Path
import json
import re
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'PRD.md'
OUT = ROOT / 'PRD.docx'
text = SOURCE.read_text(encoding='utf-8')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.7)
sec.left_margin = sec.right_margin = Inches(0.7)
styles = doc.styles
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
    style = styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.line_spacing = 1.05
styles['Title'].font.size = Pt(23)
for name, size in [('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11.5)]:
    styles[name].font.size = Pt(size)
    styles[name].paragraph_format.space_before = Pt(12)
    styles[name].paragraph_format.keep_with_next = True
doc.core_properties.title = 'Quorum product requirements'
doc.core_properties.subject = 'Product, architecture, delivery, and commercial plan'
doc.core_properties.author = 'Quorum'
doc.core_properties.keywords = 'Quorum, Agora, EchoSphere, PRD'

expected = []
links = []

def inline(p, value):
    pieces = re.split(r'(\[[^\]]+\]\(https?://[^\s)]+\)|`[^`]+`)', value)
    for part in pieces:
        match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^\s)]+)\)', part)
        if match:
            label, url = match.groups()
            rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
            h = OxmlElement('w:hyperlink')
            h.set(qn('r:id'), rid)
            r = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '175A8A'); props.append(color)
            r.append(props)
            t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)
            links.append(url)
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9)
        else:
            p.add_run(part)

def plain(value):
    value = re.sub(r'\[([^\]]+)\]\((https?://[^\s)]+)\)', r'\1', value)
    return re.sub(r'`([^`]+)`', r'\1', value)

def para(value, style=None):
    p = doc.add_paragraph(style=style)
    inline(p, value)
    expected.append(plain(value))
    return p

def table(rows):
    count = len(rows[0])
    t = doc.add_table(rows=1, cols=count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = ([3.65, 3.45] if rows[0][0] == 'Endpoint' else [2.05, 5.05]) if count == 2 else [1.35, 2.8, 2.95]
    for column, width in zip(t.columns, widths): column.width = Inches(width)
    tblpr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        b = OxmlElement('w:' + edge)
        for key, val in [('val','single'),('sz','4'),('color','D9D9D9')]: b.set(qn('w:' + key), val)
        borders.append(b)
    tblpr.append(borders)
    for i, values in enumerate(rows):
        row = t.rows[0] if i == 0 else t.add_row()
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if i == 0: pr.append(OxmlElement('w:tblHeader'))
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margin = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m = OxmlElement('w:' + side); m.set(qn('w:w'), '90'); m.set(qn('w:type'), 'dxa'); margin.append(m)
            cp.append(margin)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            inline(p, value)
            for r in p.runs:
                r.font.size = Pt(10)
                if i == 0: r.bold = True
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAF0F5'); cp.append(shade)
            expected.append(plain(value))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

lines = text.splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line.startswith('```'):
        code = []
        i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            code.append(lines[i]); i += 1
        value = '\n'.join(code)
        p = doc.add_paragraph()
        r = p.add_run(value); r.font.name = 'Consolas'; r.font.size = Pt(9)
        p.paragraph_format.line_spacing = 1.0
        expected.append(value)
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', x) for x in cells): rows.append(cells)
            i += 1
        table(rows)
        continue
    elif line.startswith('# '): para(line[2:], 'Title')
    elif line.startswith('## '): para(line[3:], 'Heading 1')
    elif line.startswith('### '): para(line[4:], 'Heading 2')
    elif line.startswith('- '): para(line[2:], 'List Bullet')
    elif re.match(r'^\d+\. ', line): para(re.sub(r'^\d+\. ', '', line), 'List Number')
    else: para(line)
    i += 1

doc.save(OUT)

NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(OUT) as z:
    xml = etree.fromstring(z.read('word/document.xml'))
    observed = []
    for p in xml.xpath('//w:body//w:p', namespaces=NS):
        parts = []
        for n in p.iter():
            if n.tag == qn('w:t'): parts.append(n.text or '')
            elif n.tag == qn('w:br'): parts.append('\n')
            elif n.tag == qn('w:tab'): parts.append('\t')
        value = ''.join(parts)
        if value: observed.append(value)
    assert expected == observed, 'Markdown/DOCX content mismatch'
    rel = etree.fromstring(z.read('word/_rels/document.xml.rels'))
    output_links = [n.get('Target') for n in rel if n.get('Type') == RT.HYPERLINK]
    assert set(links) == set(output_links), 'Hyperlink mismatch'
assert not re.search('[\u2013\u2014\u2018\u2019\u201c\u201d]', text), 'Unexpected punctuation'
report = {
    'markdown_words': len(text.split()),
    'content_blocks_verified': len(expected),
    'hyperlinks_verified': len(links),
    'unique_hyperlinks': len(set(links)),
    'docx_bytes': OUT.stat().st_size,
    'content_parity': True,
    'tables': len(doc.tables),
    'visual_review': 'pending'
}
(Path(__file__).parent / 'verification.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
(Path(__file__).parent / 'report-source.md').write_text(text, encoding='utf-8')
print(json.dumps(report, indent=2))
